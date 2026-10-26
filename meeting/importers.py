"""
Bulk import of songs into a hymn book.

Accepted input (all produce the same list of rows):

1. Spreadsheet: .csv or .xlsx with a header row. Column names may be English or
   Swahili, in any order:
       number / namba / no        (required)
       title  / kichwa / jina     (required)
       lyrics / words / maneno    (required)
       author / mwandishi         (optional)
       type   / aina              (optional: hymn, poem, reading — or wimbo, shairi, somo)
2. Pasted text: each song starts with a line "123. Title" (or "123 Title" / "123) Title").
   An optional next line "Author: …" / "Mwandishi: …". Everything until the next
   numbered line is the words (verses separated by blank lines are kept).

Validation never writes; `import_rows` writes all-or-nothing in one transaction.
"""
import csv
import io
import re

from django.db import transaction

from .models import Song

MAX_ROWS = 2000
MAX_FILE_BYTES = 5 * 1024 * 1024
MAX_LYRICS = 20000

HEADER_ALIASES = {
    'number': {'number', 'no', 'no.', 'num', 'namba', 'nambari', '#', 'hymn number', 'song number'},
    'title': {'title', 'name', 'kichwa', 'jina', 'jina la wimbo', 'song', 'hymn'},
    'lyrics': {'lyrics', 'words', 'text', 'maneno', 'mashairi', 'beti'},
    'author': {'author', 'writer', 'composer', 'mwandishi', 'mtunzi'},
    'song_type': {'type', 'song_type', 'kind', 'aina'},
}
TYPE_ALIASES = {
    'hymn': 'hymn', 'wimbo': 'hymn', 'nyimbo': 'hymn', 'song': 'hymn', '': 'hymn',
    'poem': 'poem', 'shairi': 'poem', 'spiritual poem': 'poem',
    'reading': 'reading', 'somo': 'reading', 'scripture reading': 'reading', 'scripture': 'reading',
}
SONG_START = re.compile(r'^\s*(\d{1,4})\s*[.)\-:]?\s+(\S.*)$')
AUTHOR_LINE = re.compile(r'^\s*(author|mwandishi|mtunzi|by)\s*[:\-]\s*(.+)$', re.I)


class ImportFormatError(ValueError):
    """The file as a whole couldn't be read (wrong format, missing columns…)."""


# ---------------------------------------------------------------------------
# Parsing
# ---------------------------------------------------------------------------
def _map_headers(headers):
    mapping = {}
    for i, h in enumerate(headers):
        key = str(h or '').strip().lower()
        for field, names in HEADER_ALIASES.items():
            if key in names and field not in mapping:
                mapping[field] = i
    missing = [f for f in ('number', 'title', 'lyrics') if f not in mapping]
    if missing:
        raise ImportFormatError(
            'The first row must name the columns. Missing: ' + ', '.join(missing)
            + '. Use number, title, lyrics (and optionally author, type) — or namba, kichwa, maneno.')
    return mapping


def _rows_from_table(table):
    table = [r for r in table if any(str(c or '').strip() for c in r)]
    if not table:
        raise ImportFormatError('The file is empty.')
    mapping = _map_headers(table[0])
    def cell(row, field):
        i = mapping.get(field)
        v = row[i] if i is not None and i < len(row) else ''
        if isinstance(v, float) and v.is_integer():  # Excel stores 12 as 12.0
            v = int(v)
        return '' if v is None else str(v)

    return [{'line': line_no, **{f: cell(r, f) for f in ('number', 'title', 'lyrics', 'author', 'song_type')}}
            for line_no, r in enumerate(table[1:], start=2)]


def _decode(raw):
    for enc in ('utf-8-sig', 'cp1252'):  # Excel "CSV" on Windows is often cp1252
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    raise ImportFormatError('Could not read the file’s text. Save it as “CSV UTF-8”.')


def parse_csv(raw):
    text = _decode(raw)
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=',;\t')
    except csv.Error:
        dialect = csv.excel
    return _rows_from_table(list(csv.reader(io.StringIO(text), dialect)))


def parse_xlsx(raw):
    from openpyxl import load_workbook
    try:
        wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001
        raise ImportFormatError('That doesn’t look like an Excel (.xlsx) file.') from exc
    ws = wb.worksheets[0]
    table = [list(r) for r in ws.iter_rows(values_only=True, max_row=MAX_ROWS + 50)]
    wb.close()
    return _rows_from_table(table)


def parse_text(text):
    rows, current = [], None
    for line_no, line in enumerate(text.replace('\r\n', '\n').split('\n'), start=1):
        m = SONG_START.match(line)
        if m:
            if current:
                rows.append(current)
            current = {'line': line_no, 'number': m.group(1), 'title': m.group(2).strip(), 'lyrics_lines': [],
                       'author': '', 'song_type': ''}
            continue
        if current is None:
            if line.strip():
                raise ImportFormatError(f'Line {line_no}: each song must start with its number and title, e.g. “12. Title”.')
            continue
        a = AUTHOR_LINE.match(line)
        if a and not current['lyrics_lines'] and not current['author']:
            current['author'] = a.group(2).strip()
            continue
        current['lyrics_lines'].append(line.rstrip())
    if current:
        rows.append(current)
    if not rows:
        raise ImportFormatError('No songs found. Start each song with a line like “12. Title”.')
    for r in rows:
        r['lyrics'] = '\n'.join(r.pop('lyrics_lines')).strip('\n')
    return rows


def parse_upload(upload=None, text=None):
    """Return raw rows from an uploaded file (.csv/.xlsx/.txt) or pasted text."""
    if upload is not None:
        if upload.size > MAX_FILE_BYTES:
            raise ImportFormatError('That file is over 5 MB. Split it into smaller files.')
        name = (upload.name or '').lower()
        raw = upload.read()
        if name.endswith('.xlsx'):
            return parse_xlsx(raw)
        if name.endswith('.xls'):
            raise ImportFormatError('Old .xls files aren’t supported. In Excel choose File → Save As → .xlsx or CSV.')
        if name.endswith(('.txt', '.md')):
            return parse_text(_decode(raw))
        return parse_csv(raw)
    if text and text.strip():
        return parse_text(text)
    raise ImportFormatError('Choose a file or paste some songs.')


# ---------------------------------------------------------------------------
# Validation & import
# ---------------------------------------------------------------------------
def validate_rows(book, rows, mode='skip'):
    """
    Check every row; returns a preview list:
      {line, number, title, author, song_type, lyrics_preview, status, errors}
    status: 'new' | 'replace' (number exists, mode=replace) | 'skip' (exists, mode=skip) | 'error'
    """
    if len(rows) > MAX_ROWS:
        raise ImportFormatError(f'Up to {MAX_ROWS} songs per import, please. This file has {len(rows)}.')
    existing = set(Song.objects.filter(book=book).values_list('number', flat=True))
    seen = {}
    preview = []
    for r in rows:
        errors = []
        raw_number = str(r.get('number', '')).strip().rstrip('.')
        number = None
        if not raw_number.isdigit() or int(raw_number) < 1:
            errors.append('Number must be a whole number from 1.')
        else:
            number = int(raw_number)
            if number in seen:
                errors.append(f'Number {number} appears twice in this file (also line {seen[number]}).')
            else:
                seen[number] = r['line']
        title = ' '.join(str(r.get('title', '')).split())
        if not title:
            errors.append('Title is missing.')
        elif len(title) > 255:
            errors.append('Title is longer than 255 characters.')
        lyrics = str(r.get('lyrics', '')).replace('\r\n', '\n').strip()
        if not lyrics:
            errors.append('Words are missing.')
        elif len(lyrics) > MAX_LYRICS:
            errors.append('Words are unusually long (over 20,000 characters).')
        author = ' '.join(str(r.get('author', '')).split())[:255]
        type_key = str(r.get('song_type', '')).strip().lower()
        song_type = TYPE_ALIASES.get(type_key)
        if song_type is None:
            errors.append('Type must be hymn, poem or reading (or wimbo, shairi, somo).')

        if errors:
            status = 'error'
        elif number in existing:
            status = 'replace' if mode == 'replace' else 'skip'
        else:
            status = 'new'
        preview.append({
            'line': r['line'], 'number': number if number is not None else raw_number, 'title': title,
            'author': author, 'song_type': song_type or type_key, 'lyrics': lyrics,
            'lyrics_preview': lyrics.split('\n')[0][:80], 'status': status, 'errors': errors,
        })
    return preview


def summarize(preview):
    counts = {'new': 0, 'replace': 0, 'skip': 0, 'error': 0}
    for p in preview:
        counts[p['status']] += 1
    return counts


@transaction.atomic
def import_rows(book, preview):
    """Write a validated preview. Refuses if any row has errors (all-or-nothing)."""
    if any(p['status'] == 'error' for p in preview):
        raise ImportFormatError('Fix the rows marked with errors first.')
    new = [Song(book=book, number=p['number'], title=p['title'], author=p['author'], song_type=p['song_type'],
                lyrics=p['lyrics']) for p in preview if p['status'] == 'new']
    Song.objects.bulk_create(new, batch_size=500)
    replaced = 0
    for p in preview:
        if p['status'] == 'replace':
            replaced += Song.objects.filter(book=book, number=p['number']).update(
                title=p['title'], author=p['author'], song_type=p['song_type'], lyrics=p['lyrics'])
    return {'created': len(new), 'replaced': replaced, 'skipped': sum(p['status'] == 'skip' for p in preview)}
